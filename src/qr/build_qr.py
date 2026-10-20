# -*- coding: utf-8 -*-
"""Ma QR dan toi trang So quy bao tri SSR: anh A4 de in dan bang tin, ban PDF, ban Word sua duoc, anh 16:9 de chieu Hoi nghi.
Hai kho cung mot noi dung (14/10/2026): quet ma, dua bieu tuong ra man hinh dien thoai, trong trang co, duong dan, mien tru. Nhan dien SSR."""
import base64, os, segno, cv2
from playwright.sync_api import sync_playwright
O = "/home/claude/ssr/qr"; URL = "https://jukytwo.github.io/ssr-so-quy-bao-tri/"; UPD = "14/10/2026"
q = segno.make(URL, error="q"); q.save(f"{O}/qr.png", scale=24, border=2, dark="#12301F", light="#FFFFFF")
b64 = lambda p: base64.b64encode(open(p, "rb").read()).decode()
QR = "data:image/png;base64," + b64(f"{O}/qr.png"); LOGO = "data:image/png;base64," + b64("/home/claude/ssr/crops/logo_ssr_512.png")
APPIC = "data:image/png;base64," + b64("/home/claude/ssr/icons/apple-touch-icon.png")
CSS = open("/home/claude/ssr/bvp_fonts.css").read() + open("/home/claude/ssr/fonts/cormorant.css").read()
SER = "font-family:'Cormorant Garamond',serif;font-size-adjust:.52;font-variant-numeric:lining-nums;font-weight:700;color:#F1E6C6"
TITLE = "Quỹ bảo trì chung cư: xem <em>phân tích và bộ clip</em> ngay trên điện thoại"
TITLE_TXT = "Quỹ bảo trì chung cư: xem phân tích và bộ clip ngay trên điện thoại"
INSIDE = ["Câu trả lời ngắn: tiền quỹ có mất không", "Năm vấn đề, mức độ và khung xử lý", "198 giao dịch của ba tài khoản, tra cứu được",
          "Mười câu hỏi đã gửi Ban quản trị", "Bộ clip: bản đầy đủ 22 phút, 6 phần, 11 clip dọc, tóm tắt 90 giây"]
STEPS = [("1", "Mở camera điện thoại, hướng vào mã"), ("2", "Bấm vào đường dẫn hiện ra"), ("3", "Không cần đăng nhập, không cần cài ứng dụng")]
HOME_T = "Giữ trên màn hình điện thoại như một ứng dụng"
HOME_S = "Làm một lần, sau đó chỉ cần chạm vào biểu tượng con dấu SSR để mở."
IOS = ["Mở trang bằng Safari", "Bấm biểu tượng Chia sẻ ở thanh dưới", "Chọn “Thêm vào MH chính”, bấm Thêm"]
AND = ["Mở trang bằng Chrome", "Bấm dấu ba chấm ở góc trên", "Chọn “Thêm vào màn hình chính”, bấm Thêm"]
ZALO = "Nếu quét bằng Zalo, trang mở trong Zalo: bấm dấu ba chấm, chọn “Mở bằng trình duyệt” rồi làm như trên."
DISC = "Phân tích độc lập phục vụ giám sát của cư dân, không phải văn bản của chủ đầu tư, Ban quản trị hay đơn vị quản lý vận hành. Các nhận định là dấu hiệu cần làm rõ, không phải cáo buộc."
BG = "background:radial-gradient(1200px 900px at 0% 0%,rgba(70,140,95,.55),transparent 70%),linear-gradient(160deg,#24543A,#12301F 55%,#0B1F14)"
SV_SHARE = '<svg viewBox="0 0 24 24" fill="none" stroke="#12301F" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"><path d="M8 9H6.5A1.5 1.5 0 0 0 5 10.5v9A1.5 1.5 0 0 0 6.5 21h11a1.5 1.5 0 0 0 1.5-1.5v-9A1.5 1.5 0 0 0 17.5 9H16"/><path d="M12 14V3"/><path d="M8.5 6.5 12 3l3.5 3.5"/></svg>'
SV_DOTS = '<svg viewBox="0 0 24 24" fill="#12301F"><circle cx="12" cy="5" r="2.2"/><circle cx="12" cy="12" r="2.2"/><circle cx="12" cy="19" r="2.2"/></svg>'

def card(name, glyph, steps):
    li = "".join(f"<li><i>{k}</i><span>{t}</span></li>" for k, t in zip("123", steps))
    return f'<div class="os"><div class="osh"><b>{name}</b><em class="gl">{glyph}</em></div><ol>{li}</ol></div>'

def page(K):
    """K: bo kich thuoc cho tung kho."""
    steps = "".join(f'<div><i>{n}</i><span>{t}</span></div>' for n, t in STEPS)
    inside = "".join(f"<li>{x}</li>" for x in INSIDE)
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>{CSS}
@page{{size:A4;margin:0}}*{{box-sizing:border-box;margin:0;padding:0}}
body{{width:{K['W']}px;height:{K['H']}px;font-family:'Be Vietnam Pro';color:#fff;{BG};position:relative;overflow:hidden;-webkit-print-color-adjust:exact;print-color-adjust:exact}}
.fr{{position:absolute;inset:{K['fr']}px;border:2px solid rgba(227,200,120,.55);border-radius:30px;pointer-events:none}}
.pg{{position:absolute;inset:{K['pad']}px;display:flex;flex-direction:column}}
.hd{{display:flex;gap:22px;align-items:center}}
.hd img{{width:{K['logo']}px;height:{K['logo']}px;border-radius:50%;box-shadow:0 0 0 6px rgba(255,255,255,.12)}}
.hd b{{display:block;font-size:{K['brand']}px;letter-spacing:.12em;color:#E3C878;font-weight:800}} .hd span{{font-size:{K['brand2']}px;color:rgba(255,255,255,.82)}}
h1{{{SER};font-size:{K['h1']}px;line-height:1.14;margin-top:{K['h1m']}px}} h1 em{{font-style:normal;color:#E3C878}}
.main{{display:grid;grid-template-columns:{K['cols']};gap:{K['gap']}px;align-items:center;margin-top:{K['mainm']}px}}
.qrb{{background:#fff;border-radius:26px;padding:22px;box-shadow:0 18px 50px rgba(0,0,0,.35);width:{K['qr']}px}}
.qrb img{{display:block;width:100%}}
.qside{{display:flex;flex-direction:column;gap:14px}}
.cap{{font-size:{K['cap']}px;font-weight:700;color:#E3C878;line-height:1.2}}
.st{{display:flex;flex-direction:column;gap:12px}}
.st div{{display:flex;gap:14px;align-items:center;font-size:{K['st']}px;line-height:1.3}}
.st i,.os li i{{font-style:normal;flex:none;border-radius:50%;background:#E3C878;color:#12301F;font-weight:800;display:grid;place-items:center}}
.st i{{width:{K['sti']}px;height:{K['sti']}px;font-size:{K['st']*0.8:.0f}px}}
.url{{font-size:{K['url']}px;color:#CFE3D6;margin-top:4px}} .url b{{color:#fff;font-weight:600}}
.home{{margin-top:{K['homem']}px;border:1.5px solid rgba(227,200,120,.55);background:rgba(255,255,255,.06);border-radius:22px;padding:{K['homep']}px}}
.hh{{display:flex;gap:18px;align-items:center}}
.hh img{{width:{K['appic']}px;height:{K['appic']}px;border-radius:22%;box-shadow:0 6px 18px rgba(0,0,0,.35);flex:none}}
.hh .t{{font-size:{K['ht']}px;font-weight:700;color:#E3C878;line-height:1.2}} .hh .s{{font-size:{K['hs']}px;color:#EAF2EC;margin-top:4px;line-height:1.35}}
.oss{{display:grid;grid-template-columns:1fr 1fr;gap:16px;margin-top:14px}}
.os{{background:rgba(0,0,0,.18);border-radius:16px;padding:14px 18px}}
.osh{{display:flex;align-items:center;justify-content:space-between}}
.osh b{{font-size:{K['osb']}px;color:#fff}}
.gl{{display:grid;place-items:center;width:{K['gl']}px;height:{K['gl']}px;background:#F1E6C6;border-radius:10px}} .gl svg{{width:66%;height:66%}}
.os ol{{list-style:none;display:flex;flex-direction:column;gap:7px;margin-top:8px}}
.os li{{display:flex;gap:10px;align-items:flex-start;font-size:{K['osl']}px;line-height:1.35;color:#EAF2EC}}
.os li i{{width:{K['osi']}px;height:{K['osi']}px;font-size:{K['osl']*0.78:.0f}px;margin-top:1px}}
.zl{{font-size:{K['zl']}px;color:rgba(255,255,255,.78);margin-top:12px;line-height:1.4}}
.in{{margin-top:{K['inm']}px}}
.in b{{font-size:{K['inb']}px;letter-spacing:.12em;color:#E3C878;font-weight:700}}
.in ul{{list-style:none;display:grid;grid-template-columns:1fr 1fr;gap:5px 30px;margin-top:8px}}
.in li{{font-size:{K['inl']}px;line-height:1.38;color:#EAF2EC;padding-left:20px;position:relative}} .in li::before{{content:"";position:absolute;left:0;top:{K['inl']*0.55:.0f}px;width:8px;height:8px;border-radius:50%;background:#E3C878}}
.in li:last-child{{grid-column:1/-1}}
.rc{{display:contents}}
.ft{{margin-top:auto;font-size:{K['ft']}px;line-height:1.45;color:rgba(255,255,255,.68)}} .ft u{{text-decoration:none;color:#CFE3D6}}
{K.get('extra', '')}
</style></head><body><div class="fr"></div><div class="pg">
<div class="hd"><img src="{LOGO}"><div><b>SỔ QUỸ BẢO TRÌ SSR</b><span>Chung cư Nam Sài Gòn Residences</span></div></div>
<h1>{TITLE}</h1>
<div class="main"><div class="qrb"><img src="{QR}"></div>
<div class="qside"><div class="cap">Quét mã bằng camera điện thoại</div><div class="st">{steps}</div>
<div class="url">Hoặc gõ: <b>{URL.replace('https://', '')}</b></div></div></div>
<div class="rc"><div class="home"><div class="hh"><img src="{APPIC}"><div><div class="t">{HOME_T}</div><div class="s">{HOME_S}</div></div></div>
<div class="oss">{card('iPhone', SV_SHARE, IOS)}{card('Android', SV_DOTS, AND)}</div><div class="zl">{ZALO}</div></div>
<div class="in"><b>TRONG TRANG CÓ</b><ul>{inside}</ul></div></div>
<div class="ft">Đường dẫn: <u>{URL.replace('https://', '')}</u> · Cập nhật {UPD}<br>{DISC}</div>
</div></body></html>"""

KA = dict(W=1240, H=1754, fr=28, pad=84, logo=112, brand=30, brand2=23, h1=62, h1m=26, cols="460px 1fr", gap=40, mainm=30, qr=460,
          cap=31, st=24, sti=44, url=21, homem=30, homep=22, appic=78, ht=27, hs=20, osb=24, gl=40, osl=20, osi=28, zl=18,
          inm=24, inb=20, inl=20, ft=16)
# kho ngang: hai cot, cot trai la QR va buoc quet, cot phai la man hinh dien thoai va noi dung
KW = dict(W=1920, H=1080, fr=26, pad=72, logo=84, brand=26, brand2=20, h1=52, h1m=14, cols="440px 1fr", gap=32, mainm=22, qr=440,
          cap=27, st=22, sti=40, url=18, homem=0, homep=26, appic=80, ht=29, hs=20, osb=25, gl=42, osl=22, osi=31, zl=18,
          inm=30, inb=19, inl=21, ft=14,
          extra=""".pg{display:grid;grid-template-columns:830px 1fr;grid-template-rows:auto auto 1fr auto;column-gap:56px}
.hd{grid-column:1;grid-row:1} h1{grid-column:1;grid-row:2} .main{grid-column:1;grid-row:3;align-self:start;grid-template-columns:440px 1fr}
.url b{display:block;white-space:nowrap}
.oss{gap:18px;margin-top:20px} .os{padding:18px 22px} .os ol{gap:12px;margin-top:12px} .in ul{gap:10px 30px;margin-top:12px}
.rc{display:flex!important;flex-direction:column;grid-column:2;grid-row:1/4;align-self:start} .ft{grid-column:1/-1;grid-row:4;margin-top:12px}""")
A4 = page(KA); W = page(KW)
open(f"{O}/a4.html", "w").write(A4); open(f"{O}/w.html", "w").write(W)
NA = "Mã QR - Sổ quỹ bảo trì SSR (A4 in dán bảng tin)"; NW = "Mã QR - Sổ quỹ bảo trì SSR (chiếu Hội nghị 16x9)"
CHK = """(()=>{const H=document.body.clientHeight,W=document.body.clientWidth,els=[...document.querySelectorAll('.hd,h1,.main,.home,.in,.ft')].map(e=>e.getBoundingClientRect());
const ov=[];for(let i=0;i<els.length;i++)for(let j=i+1;j<els.length;j++){const a=els[i],b=els[j];if(a.left<b.right&&b.left<a.right&&a.top<b.bottom&&b.top<a.bottom)ov.push(i+'-'+j)}
const pg=document.querySelector('.pg');return {overlap:ov,overflow:pg.scrollHeight>pg.clientHeight+1||pg.scrollWidth>pg.clientWidth+1,
gap_before_ft:Math.round(els[5].top-Math.max(els[3].bottom,els[4].bottom)),bottom:Math.round(els[5].bottom),H}})()"""
with sync_playwright() as p:
    b = p.chromium.launch()
    pg = b.new_page(viewport={"width": 1240, "height": 1754}, device_scale_factor=2); pg.goto(f"file://{O}/a4.html"); pg.wait_for_timeout(800)
    print("A4:", pg.evaluate(CHK))
    pg.screenshot(path=f"{O}/{NA}.png"); pg.pdf(path=f"{O}/{NA}.pdf", width="210mm", height="297mm", print_background=True, scale=1/(1240/793.7))
    pw = b.new_page(viewport={"width": 1920, "height": 1080}); pw.goto(f"file://{O}/w.html"); pw.wait_for_timeout(800)
    print("16:9:", pw.evaluate(CHK))
    pw.screenshot(path=f"{O}/{NW}.png"); b.close()
det = cv2.QRCodeDetector()
for f in [f"{O}/{NA}.png", f"{O}/{NW}.png"]:
    im = cv2.imread(f); im = cv2.resize(im, (im.shape[1] // 2, im.shape[0] // 2)); v, *_ = det.detectAndDecode(im); print("quet:", v == URL, os.path.basename(f))
# ban Word sua duoc
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
d = Document(); s = d.sections[0]; s.page_width, s.page_height = Mm(210), Mm(297); s.left_margin = s.right_margin = Mm(20); s.top_margin = s.bottom_margin = Mm(14)
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12.5)
st.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Times New Roman")
G = RGBColor(0x1E, 0x46, 0x30); GO = RGBColor(0xA8, 0x83, 0x2E); GR = RGBColor(0x55, 0x55, 0x55)
def para(t, size=12.5, bold=False, color=None, align=None, after=4):
    p = d.add_paragraph(); r = p.add_run(t); r.font.size = Pt(size); r.bold = bold
    if color: r.font.color.rgb = color
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(after); return p
C = WD_ALIGN_PARAGRAPH.CENTER
p = d.add_paragraph(); p.alignment = C; p.add_run().add_picture("/home/claude/ssr/crops/logo_ssr_512.png", width=Mm(18))
para("SỔ QUỸ BẢO TRÌ SSR", 14, True, GO, C, 0); para("Chung cư Nam Sài Gòn Residences", 11.5, False, None, C, 8)
para(TITLE_TXT, 20, True, G, C, 6)
p = d.add_paragraph(); p.alignment = C; p.add_run().add_picture(f"{O}/qr.png", width=Mm(62))
para("Quét mã bằng camera điện thoại", 14, True, GO, C, 4)
for n, t in STEPS: para(f"{n}. {t}", 12.5, False, None, C, 1)
para("Hoặc gõ: " + URL.replace("https://", ""), 11.5, False, GR, C, 8)
para(HOME_T.upper(), 12, True, G, None, 1); para(HOME_S, 12, False, None, None, 3)
tb = d.add_table(rows=1, cols=2); tb.style = "Table Grid"
for cell, (nm, ss) in zip(tb.rows[0].cells, [("iPhone", IOS), ("Android", AND)]):
    cell.paragraphs[0].add_run(nm).bold = True
    for k, t in enumerate(ss, 1):
        cp = cell.add_paragraph(f"{k}. {t}"); cp.paragraph_format.space_after = Pt(1)
para(ZALO, 11, False, GR, None, 8)
para("TRONG TRANG CÓ", 12, True, G, None, 2)
for x in INSIDE: para("•  " + x, 12, False, None, None, 1)
para("", 4); para("Đường dẫn: " + URL.replace("https://", "") + " · Cập nhật " + UPD, 10, False, GR, None, 2)
para(DISC, 10, False, GR, None, 0)
d.save(f"{O}/{NA}.docx"); print("docx ok")
